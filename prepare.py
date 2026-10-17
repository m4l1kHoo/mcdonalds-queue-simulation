"""Read historical files as data, never execute notebooks or recalculate workbooks."""
from __future__ import annotations

import argparse
import ast
import csv
from datetime import datetime, time, timedelta
import hashlib
import json
import math
from pathlib import Path
import re

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
FIELDS = ['sample_id', 'dataset', 'session', 'sequence', 'source_file', 'source_sheet',
          'source_cells', 'raw_value', 'raw_unit', 'value', 'unit', 'status', 'note']


def save_json(path, value):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False) + '\n')


def save_csv(path, rows, fields=None):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open('w', newline='', encoding='utf-8') as f:
        writer = csv.DictWriter(f, fieldnames=fields or list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def read_csv(path):
    with Path(path).open(newline='', encoding='utf-8') as f:
        return list(csv.DictReader(f))


def decimal_minutes(value):
    result = float(value) * 60
    if not math.isfinite(result) or result < 0:
        raise ValueError('Duration must be finite and nonnegative')
    return result


def clock_seconds(value):
    if isinstance(value, datetime):
        value = value.time()
    if isinstance(value, time):
        return value.hour * 3600 + value.minute * 60 + value.second + value.microsecond / 1e6
    raise ValueError('Expected an Excel clock, not a decimal duration')


def duration_between(start, end):
    # No silent midnight roll-over or date correction.
    if isinstance(start, datetime) and isinstance(end, datetime):
        seconds = (end - start).total_seconds()
        status = 'date_conflict' if start.date() != end.date() else 'valid'
    else:
        seconds = clock_seconds(end) - clock_seconds(start)
        status = 'valid'
    if seconds < 0:
        status = 'negative_interval'
    return seconds, status


def explicit_duration(text):
    match = re.fullmatch(r'(\d+)min(\d+)s', str(text))
    if not match or int(match[2]) >= 60:
        raise ValueError('Expected minutes and seconds with seconds < 60')
    return int(match[1]) * 60 + int(match[2])


def literal_arrays(path):
    """Extract literal list assignments; no imports, eval, exec or saved output execution."""
    arrays = {}
    nb = json.loads(Path(path).read_text())
    for cell_number, cell in enumerate(nb['cells'], 1):
        if cell['cell_type'] != 'code':
            continue
        for node in ast.parse(''.join(cell['source'])).body:
            if isinstance(node, ast.Assign) and isinstance(node.value, (ast.List, ast.Tuple)):
                for target in node.targets:
                    if isinstance(target, ast.Name):
                        try:
                            value = ast.literal_eval(node.value)
                        except (ValueError, TypeError):
                            continue  # e.g. a tuple of fitted parameter names, not literal data
                        arrays[target.id] = (cell_number, value)
    return arrays


class Preparation:
    def __init__(self, source):
        self.source = Path(source)
        self.rows, self.arrivals, self.inventory, self.reconciliation = [], [], [], []

    def add(self, dataset, file, sheet, cells, raw, raw_unit, value, *, status='valid',
            unit='seconds', note='', session=None):
        sequence = sum(r['dataset'] == dataset for r in self.rows) + 1
        if isinstance(value, float):
            value = round(value, 9)
        self.rows.append(dict(zip(FIELDS, [f'{dataset}:{sequence:03}', dataset,
            session or file, sequence, file, sheet, cells, str(raw), raw_unit, value,
            unit, status, note])))

    def workbook(self, name):
        return load_workbook(self.source / 'Collected_data' / name, read_only=True, data_only=True)

    def check_sources(self):
        manifest = json.loads((ROOT / 'docs/source_manifest.json').read_text())
        for item in manifest:
            if item['copied'] and Path(item['source']).suffix in ['.xlsx', '.ipynb']:
                p = self.source / item['source']
                if hashlib.sha256(p.read_bytes()).hexdigest() != item['sha256']:
                    raise ValueError(f'Source hash mismatch: {item["source"]}')
        for p in sorted((self.source / 'Collected_data').glob('*.xlsx')):
            w = load_workbook(p, read_only=True, data_only=False)
            for s in w:
                cells = [c for row in s for c in row if c.value is not None]
                self.inventory.append({'file':p.name, 'sheet':s.title,
                    'nonempty_cells':len(cells), 'formula_cells':[c.coordinate for c in cells if c.data_type == 'f'],
                    'handling':'Field mappings in prepare.py; observer names, example rows, order IDs and free text are not exported.'})
            w.close()

    def orders(self):
        name = '0812_1200_1300_订单系统.xlsx'
        w = self.workbook(name); s = w['Sheet1']
        times = []
        for row in range(2, 68):
            a, b = s[f'B{row}'].value, s[f'C{row}'].value
            value, status = duration_between(a, b)
            self.add('order_completion', name, s.title, f'B{row}:C{row}', f'{a}|{b}', 'clock', value,
                     status=status, note='Order arrival to completion; includes all intervening delay, not isolated kitchen service.')
            prefix = str(s[f'A{row}'].value)[:2]
            self.add('order_prefix35', name, s.title, f'A{row}', 'prefix category only', 'category', int(prefix == '35'),
                     unit='binary', note='Prefix 35 versus 39; phone/kiosk meaning is a historical hypothesis, not established by column label.')
            times.append(clock_seconds(a))
        for i in range(1, len(times)):
            value = times[i] - times[i-1]
            self.add('order_interarrival', name, s.title, f'B{i+1}:B{i+2}', 'adjacent clock difference', 'clock', value,
                     status='valid' if value >= 0 else 'negative_interval', note='Order arrivals; not customer door arrivals.')
        w.close()

    def dining(self):
        name = '用餐时间分布数据表.xlsx'; w = self.workbook(name); s = w['Sheet1']
        for row in range(2, 17):
            start, end, text = [s[f'{col}{row}'].value for col in 'BCD']
            # Excel day fractions have high precision. Round to observed second resolution.
            value = round((end - start) * 86400)
            stated = explicit_duration(text)
            self.add('dining_duration', name, s.title, f'B{row}:D{row}', f'{start}|{end}|{text}',
                     'Excel days; explicit min/s', value, note=f'Timestamp difference; stated text differs by {stated-value} seconds. Observation can describe a table/group, not necessarily one person.')
        w.close()

    def pickup(self):
        name = 'pick_immediate.xlsx'; w = self.workbook(name); s = w['Sheet1']
        for row in range(2, 82):
            value = s[f'B{row}'].value
            self.add('pick_leave_flag', name, s.title, f'B{row}', value, 'binary', value, unit='binary',
                     note='Header pick - leave does not unambiguously identify before/after ordering. Row 82 is a summary, excluded.')
        w.close()
        name = '叫号系统_2.xlsx'; w = self.workbook(name); s = w['工作表 1']
        for row in range(5, 67):
            start, end = s[f'B{row}'].value, s[f'C{row}'].value
            if isinstance(start, (int, float)) and isinstance(end, (int, float)):
                value = (end-start)*60
                status = 'valid' if value >= 0 else 'negative_interval'
            else:
                value, status = '', 'missing_interval'
            self.add('call2_elapsed', name, s.title, f'B{row}:C{row}', f'{start}|{end}',
                     'decimal minutes', value, status=status,
                     note='Columns have no event labels. Initial three pairs decrease; not repaired. Valid numeric interval does not prove same event as order completion.')
        w.close()

    def observation_forms(self):
        name = '叫号系统_1.xlsx'; w = self.workbook(name); s = w['Sheet2']
        for row in range(10, 42):
            for col, dataset in [('E','kiosk_queue_evening'),('F','kiosk_service_evening')]:
                value = s[f'{col}{row}'].value
                if isinstance(value,(int,float)):
                    self.add(dataset,name,s.title,f'{col}{row}',value,'seconds (header)',value)
            a,b = s[f'H{row}'].value,s[f'I{row}'].value
            if a is None and b is None: continue
            if a is None or b is None: value,status='','missing_interval'
            else: value,status=duration_between(a,b)
            self.add('call_to_pickup_evening',name,s.title,f'H{row}:I{row}',f'{a}|{b}',
                     'datetime',value,status=status,note='Actual dates retained; inconsistent dates are not wrapped modulo 24 hours.')
        w.close()
        name = '周六中午.xlsx'; w=self.workbook(name);s=w['Sheet2']; by_type={}
        origin = clock_seconds(s['B10'].value)
        for row in range(10,66):
            arrival = s[f'B{row}'].value
            group_size = s[f'C{row}'].value
            note = str(s[f'O{row}'].value or '')
            kind = 'courier' if '外卖员' in note else 'walkaway' if '看了一眼走了' in note else 'customer'
            offset = clock_seconds(arrival)-origin
            self.arrivals.append({'arrival_id':f'saturday:{row}', 'source_file':name, 'source_sheet':s.title,
                'source_cells':f'B{row}:O{row}', 'arrival_seconds':offset,'kind':kind,
                'group_size':group_size if group_size is not None else '',
                'kiosk_observed':int(s[f'D{row}'].value in [1,2] or isinstance(s[f'G{row}'].value,(int,float))),
                'dine_in_observed':s[f'N{row}'].value if s[f'N{row}'].value in [0,1] else '',
                'status':'observed_group; group size may be missing; no assumption exported'})
            if kind != 'walkaway':
                if kind in by_type:
                    prev_t,prev_row=by_type[kind]
                    self.add(f'{kind}_group_interarrival',name,s.title,f'B{prev_row},B{row}',
                             'consecutive arrivals in this class','clock',offset-prev_t,
                             note='Groups, not individuals. Minute-rounded observations allow zero intervals.')
                by_type[kind]=(offset,row)
            for col, dataset in [('F','kiosk_queue_saturday'),('G','kiosk_service_saturday')]:
                value=s[f'{col}{row}'].value
                if isinstance(value,(int,float)):
                    self.add(dataset,name,s.title,f'{col}{row}',value,'seconds (header)',value)
        w.close()
        name='周日中午_就座时间.xlsx';w=self.workbook(name);s=w['Sheet2']
        for row in range(10,18):
            a,b=s[f'B{row}'].value,s[f'D{row}'].value
            value,status=duration_between(a,b)
            self.add('pickup_to_leave_sunday',name,s.title,f'B{row}:E{row}',f'{a}|{b}','clock',value,
                     status=status,note='Starts at pickup, not confirmed seat acquisition; not pooled with dining duration.')
        w.close()

    def notebooks(self):
        mappings=[('order_machine_operating_time.ipynb','data','kiosk_notebook','seconds'),
                  ('delay_time_to_seat.ipynb','data','seat_delay_notebook','seconds'),
                  ('Before_Or_After_Arrival.ipynb','data','after_arrival_notebook','binary'),
                  ('Dine_in_or_takeaway.ipynb','data','dine_in_notebook','binary'),
                  ('order_completion_time.ipynb','process_time_2','completion_notebook_part2','seconds')]
        for name,key,dataset,unit in mappings:
            cell,values=literal_arrays(self.source/'Input_Analysis'/name)[key]
            for i,value in enumerate(values,1):
                self.add(dataset,name,f'cell {cell}',f'{key}[{i-1}]',value,
                         'binary' if unit=='binary' else 'decimal minutes (user confirmed)',
                         value if unit=='binary' else decimal_minutes(value),unit=unit,
                         note='Historical literal array; not a new independent observation; worksheet lineage remains incomplete.')
        arrays=literal_arrays(self.source/'Input_Analysis/order_completion_time.ipynb')
        old_start=arrays['order_arrival_time'][1];old_end=arrays['order_completion_time'][1]
        w=self.workbook('0812_1200_1300_订单系统.xlsx');s=w['Sheet1']
        self.reconciliation.append({'check':'completion notebook timestamps versus workbook',
            'n':len(old_start),'exact_matches':sum(str(s[f'B{i+2}'].value)==a and str(s[f'C{i+2}'].value)==b
              for i,(a,b) in enumerate(zip(old_start,old_end))), 'handling':'Use workbook once; do not double count notebook copy.'})
        w.close()
        old=arrays['process_time_2'][1]
        candidates=[r for r in self.rows if r['dataset']=='call2_elapsed' and r['value']!='']
        if len(old)!=len(candidates): raise ValueError('Unexpected historical array length')
        for i,(value,row) in enumerate(zip(old,candidates)):
            if abs(decimal_minutes(value)-row['value'])>1e-6:
                self.reconciliation.append({'check':'completion second array versus call2 interval',
                    'notebook_index':i, 'source_cells':row['source_cells'],
                    'notebook_seconds':decimal_minutes(value),'worksheet_seconds':row['value'],
                    'handling':'Keep separate; no inferred correction or pooling.'})
        before=[r['value'] for r in self.rows if r['dataset']=='after_arrival_notebook']
        pickup=[r['value'] for r in self.rows if r['dataset']=='pick_leave_flag']
        self.reconciliation.append({'check':'before/after array versus pick - leave worksheet',
            'complements_at_every_index':all(a+b==1 for a,b in zip(before,pickup)),
            'n':len(before),'handling':'Numerical complement does not establish identical event semantics.'})

    def run(self, output):
        self.check_sources();self.orders();self.dining();self.pickup();self.observation_forms();self.notebooks()
        output=Path(output)
        save_csv(output/'observations.csv',self.rows,FIELDS)
        save_csv(output/'arrivals.csv',self.arrivals)
        save_json(output/'source_structure.json',self.inventory)
        save_json(output/'reconciliation.json',self.reconciliation)
        datasets={}
        for r in self.rows:
            d=datasets.setdefault(r['dataset'],{'n':0,'valid':0,'unit':r['unit'],'source_file':r['source_file'],'statuses':{}})
            d['n']+=1;d['valid']+=r['status']=='valid';d['statuses'][r['status']]=d['statuses'].get(r['status'],0)+1
        save_json(output/'datasets.json',datasets)
        manifest={p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in sorted(output.iterdir()) if p.is_file() and p.suffix in ['.csv','.json'] and p.name!='manifest.json'}
        save_json(output/'manifest.json',manifest)
        return datasets


if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source',type=Path,default=ROOT/'private_originals')
    parser.add_argument('--output',type=Path,default=ROOT/'data')
    args=parser.parse_args()
    print(json.dumps(Preparation(args.source).run(args.output),ensure_ascii=False,indent=2))
